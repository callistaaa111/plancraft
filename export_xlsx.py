"""Excel export: item summary, quantities by page, the PANEL circuit summary and the page list."""
from __future__ import annotations

import io
import re
from collections import Counter, defaultdict
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import pricing
from measure import item_quantities, unit_info

THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="D9E1F2")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
BOLD = Font(bold=True)
NUM = "#,##0.00"
TYPE_LABEL = {"linear": "Linear", "area": "Area", "count": "Count"}

# Circuit names: optional 2-letter site prefix, panel, circuit number, then anything else (e.g. AB-LP1-3 (NO HOMERUN)).
CIRCUIT_NAME = re.compile(r"^(?:[A-Za-z]{2}-)?(?P<panel>.+?)-(?P<num>\d+(?:\s*&\s*\d+)?)\s*(?P<rest>.*)$")


def walk_items(nodes, path=()):
    """Yield (folder_path, item) in tree order."""
    for n in nodes:
        if n.get("kind") == "folder":
            yield from walk_items(n.get("children") or [], path + (n.get("name") or "Folder",))
        else:
            yield " / ".join(path), n


def _header(ws, row, titles, widths=None):
    for c, t in enumerate(titles, start=1):
        cell = ws.cell(row, c, t)
        cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w


def _title(ws, project, subtitle):
    ws["A1"] = project.get("name") or "Takeoff"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = subtitle
    ws["A2"].font = Font(italic=True, color="666666")


def _fmt(v):
    return f"{v:,.2f}"


def build(project: dict) -> bytes:
    units = project.get("units") or "m"
    pages = {p["id"]: p for p in project.get("pages") or []}
    ours = set()   # cells holding formulas we wrote; any other text starting with "=" is the user's and stays text

    def formula(ws, r, c, f):
        ours.add((ws.title, r, c))
        return ws.cell(r, c, f)

    page_order = {p["id"]: i for i, p in enumerate(project.get("pages") or [])}
    rows = []
    for folder, item in walk_items(project.get("tree") or []):
        rows.append((folder, item, item_quantities(item, pages, units)))
    stamp = datetime.now().strftime("%Y-%m-%d %H:%M")

    wb = Workbook()

    # ---- Summary ------------------------------------------------------------------------------------
    ws = wb.active
    ws.title = "Summary"
    _title(ws, project, f"Takeoff summary  ·  measured in {unit_info(units)['len']}  ·  exported {stamp}")
    head = ["FOLDER", "ITEM", "TYPE", "DESCRIPTION", "SPEC", "QTY", "UNIT", "ALLOW %", "QTY + ALLOW", "NOTES"]
    _header(ws, 4, head, [22, 30, 8, 26, 26, 12, 7, 9, 12, 46])
    r = 5
    for folder, item, q in rows:
        props = item.get("props") or {}
        allow = props.get("allowance") or 0
        notes = [f"{label} {_fmt(v)} {u}" for label, v, u in q["extra"]]
        if q["unscaled"]:
            notes.append(f"{q['unscaled']} section(s) on pages without a scale are not counted")
        values = [folder, item.get("name"), TYPE_LABEL.get(item.get("type"), item.get("type")),
                  props.get("description") or "", props.get("spec") or "",
                  round(q["qty"], 4), q["unit"], allow, None, "; ".join(notes)]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(r, c, v)
            cell.border = BORDER
        formula(ws, r, 9, f"=IF(H{r}>0,ROUNDUP(F{r}*(1+H{r}/100),0),F{r})")
        ws.cell(r, 6).number_format = "0" if item.get("type") == "count" else NUM
        ws.cell(r, 9).number_format = "0" if item.get("type") == "count" else NUM
        if q["unscaled"]:
            for c in (6, 10):
                ws.cell(r, c).fill = WARN_FILL
        r += 1
    ws.freeze_panes = "A5"
    if r > 5:
        ws.auto_filter.ref = f"A4:J{r - 1}"

    # ---- By page ------------------------------------------------------------------------------------
    wp = wb.create_sheet("By Page")
    _title(wp, project, "Quantities on each page")
    _header(wp, 4, ["PAGE", "FOLDER", "ITEM", "TYPE", "QTY", "UNIT"], [36, 22, 30, 8, 12, 7])
    entries = []
    for ti, (folder, item, q) in enumerate(rows):
        for pid, v in q["per_page"].items():
            entries.append((page_order.get(pid, 1e9), ti, pages[pid].get("name"), folder, item, v, q["unit"]))
    entries.sort(key=lambda e: (e[0], e[1]))
    r = 5
    for _, _, pname, folder, item, v, unit in entries:
        for c, val in enumerate([pname, folder, item.get("name"), TYPE_LABEL.get(item.get("type")), round(v, 4), unit],
                                start=1):
            wp.cell(r, c, val).border = BORDER
        wp.cell(r, 5).number_format = "0" if item.get("type") == "count" else NUM
        r += 1
    wp.freeze_panes = "A5"
    if r > 5:
        wp.auto_filter.ref = f"A4:F{r - 1}"

    # ---- Circuits (PANEL layout) --------------------------------------------------------------------
    circuits = defaultdict(dict)       # panel -> num -> [lm, desc, spec, allowance, flags]
    panel_order = []
    for _folder, item, q in rows:
        if item.get("type") != "linear":
            continue
        m = CIRCUIT_NAME.match((item.get("name") or "").strip())
        if not m:
            continue
        panel = m.group("panel").upper()
        num = re.sub(r"\s+", "", m.group("num"))
        rest = m.group("rest").strip()
        props = item.get("props") or {}
        if panel not in circuits:
            panel_order.append(panel)
        c = circuits[panel].setdefault(num, [0.0, "", "", props.get("allowance") or 0, set()])
        c[0] += q["qty"]
        c[1] = c[1] or (props.get("description") or "")
        c[2] = c[2] or (props.get("spec") or "")
        if rest:
            c[4].add(rest.strip("() ").upper())
        if q["unscaled"]:
            c[4].add("PAGE NOT SCALED")
    if circuits:
        wc = wb.create_sheet("Circuits")
        wc.merge_cells("A1:H1")
        wc["A1"] = project.get("name") or "Takeoff"
        wc["A1"].font = Font(bold=True, size=14)
        wc["A1"].alignment = Alignment(horizontal="center")
        allowances = Counter(c[3] for p in circuits.values() for c in p.values())
        common = allowances.most_common(1)[0][0] if allowances else 0
        row = 2
        for panel in panel_order:
            wc.cell(row, 1, "PANEL:").font = BOLD
            wc.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
            cell = wc.cell(row, 2, panel)
            cell.font, cell.alignment = BOLD, Alignment(horizontal="center")
            row += 1
            for col, text in enumerate(("CKT NO", "DESCRIPTION", "LM", f"LM (*{common:g}%)", "TYPE OF WIRE"), start=1):
                cell = wc.cell(row, col, text)
                cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, BORDER
                cell.alignment = Alignment(horizontal="center")
            row += 1

            def key(n):
                m = re.match(r"\d+", n)
                return int(m.group()) if m else 0
            for num in sorted(circuits[panel], key=key):
                lm, desc, spec, allow, flags = circuits[panel][num]
                if flags:
                    desc = f"{desc} ({', '.join(sorted(flags))})".strip()
                wc.cell(row, 1, int(num) if num.isdigit() else num)
                wc.cell(row, 2, desc)
                wc.cell(row, 3, round(lm, 1))
                formula(wc, row, 4, f"=ROUNDUP(C{row}*{1 + float(allow) / 100:g},0)")
                wc.cell(row, 5, spec)
                for col in range(1, 6):
                    wc.cell(row, col).border = BORDER
                if flags:
                    wc.cell(row, 3).fill = WARN_FILL
                row += 1
            row += 1
        # OTHERS: count items
        others = defaultdict(float)
        for _folder, item, q in rows:
            if item.get("type") == "count" and q["qty"]:
                others[(item.get("name") or "").upper()] += q["qty"]
        wc.merge_cells("G2:H2")
        wc["G2"] = "OTHERS"
        wc["G2"].font = BOLD
        wc["G2"].alignment = Alignment(horizontal="center")
        for col, text in ((7, "DESCRIPTION"), (8, "QTY")):
            cell = wc.cell(3, col, text)
            cell.font, cell.fill, cell.border = BOLD, HEAD_FILL, BORDER
        for i, (name, qty) in enumerate(others.items(), start=4):
            wc.cell(i, 7, name).border = BORDER
            wc.cell(i, 8, int(qty)).border = BORDER
        for col, w in zip("ABCDEFGH", (9, 28, 9, 12, 30, 3, 34, 8), strict=True):
            wc.column_dimensions[col].width = w

    # ---- Bill of materials and costs ----------------------------------------------------------------
    cost = pricing.costing(project)
    if cost["bom"]:
        cur = (project.get("settings") or {}).get("currency", "\u20b1")
        money = f'"{cur}"#,##0.00' if cur else "#,##0.00"
        wb_ = wb.create_sheet("Bill of Materials")
        _title(wb_, project, f"Bill of materials  ·  from the assemblies on each item  ·  exported {stamp}")
        _header(wb_, 4, ["CATEGORY", "MATERIAL", "UNIT", "QTY", "UNIT PRICE", "AMOUNT"], [20, 44, 8, 12, 14, 16])
        r = 5
        for b in cost["bom"]:
            m = b["material"]
            vals = [m.get("category") or "", m.get("name"), m.get("unit") or "", round(b["qty"], 4), pricing.num(m.get("price"))]
            for c, v in enumerate(vals, start=1):
                wb_.cell(r, c, v).border = BORDER
            formula(wb_, r, 6, f"=D{r}*E{r}").border = BORDER
            wb_.cell(r, 4).number_format = "0" if m.get("whole") else NUM
            wb_.cell(r, 5).number_format = money
            wb_.cell(r, 6).number_format = money
            r += 1
        last = r - 1
        rows_ = [("Subtotal", f"=SUM(F5:F{last})"),
                 (f"Markup {cost['markup']:g}%", f"=F{r}*{cost['markup'] / 100:g}"),
                 ("TOTAL", f"=F{r}+F{r + 1}")]
        for i, (label, f) in enumerate(rows_):
            wb_.cell(r + i, 5, label).font = BOLD
            c = formula(wb_, r + i, 6, f)
            c.font, c.number_format, c.border = BOLD, money, BORDER
        wb_.freeze_panes = "A5"

        wi = wb.create_sheet("Materials by Item")
        _title(wi, project, "Materials for each takeoff item (before rounding to whole units)")
        _header(wi, 4, ["FOLDER", "ITEM", "ASSEMBLY", "MATERIAL", "QTY", "UNIT", "UNIT PRICE", "AMOUNT"], [22, 28, 26, 40, 12, 8, 14, 16])
        r = 5
        for ln in cost["lines"]:
            m = ln["material"]
            vals = [ln["folder"], ln["item"].get("name"), ln["assembly"].get("name"), m.get("name"), round(ln["qty"], 4),
                    m.get("unit") or "", pricing.num(m.get("price"))]
            for c, v in enumerate(vals, start=1):
                wi.cell(r, c, v).border = BORDER
            formula(wi, r, 8, f"=E{r}*G{r}").border = BORDER
            wi.cell(r, 5).number_format = NUM
            wi.cell(r, 7).number_format = money
            wi.cell(r, 8).number_format = money
            r += 1
        wi.freeze_panes = "A5"
        if r > 5:
            wi.auto_filter.ref = f"A4:H{r - 1}"

    # ---- Notes (plan notes are not takeoff) ---------------------------------------------------------
    notes = [m for m in project.get("marks") or [] if m.get("kind") == "note" and (m.get("text") or "").strip()]
    if notes:
        wn = wb.create_sheet("Notes")
        _title(wn, project, "Notes written on the plans")
        _header(wn, 4, ["PAGE", "NOTE"], [36, 90])
        notes.sort(key=lambda m: page_order.get(m.get("page"), 1e9))
        for i, m in enumerate(notes, start=5):
            wn.cell(i, 1, (pages.get(m.get("page")) or {}).get("name", "")).border = BORDER
            c = wn.cell(i, 2, m["text"].strip())
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
        wn.freeze_panes = "A5"

    # ---- Pages --------------------------------------------------------------------------------------
    wpg = wb.create_sheet("Pages")
    _title(wpg, project, "Pages and their scales")
    _header(wpg, 4, ["#", "PAGE", "GROUP", "SCALE", "PX PER M", "DPI", "WIDTH PX", "HEIGHT PX", "SOURCE", "SECTIONS"],
            [5, 36, 22, 22, 10, 8, 10, 10, 34, 10])
    sec_count = Counter()
    for _, item, _ in rows:
        for s in item.get("sections") or []:
            sec_count[s.get("page")] += 1
    for i, p in enumerate(project.get("pages") or [], start=1):
        sc = p.get("scale") or {}
        label = sc.get("label") or ""
        if sc.get("auto"):
            label += " (read from drawing, not checked)"
        values = [i, p.get("name"), p.get("group") or "", label if sc.get("pxPerM") else "NOT SET",
                  round(sc["pxPerM"], 4) if sc.get("pxPerM") else None, p.get("dpi"), p.get("width"), p.get("height"),
                  p.get("source") or "", sec_count.get(p["id"], 0)]
        for c, v in enumerate(values, start=1):
            wpg.cell(i + 4, c, v).border = BORDER
        if not sc.get("pxPerM") and sec_count.get(p["id"]):
            wpg.cell(i + 4, 4).fill = WARN_FILL
    wpg.freeze_panes = "A5"

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" and (ws.title, cell.row, cell.column) not in ours:
                    cell.data_type = "s"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
